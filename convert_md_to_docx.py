import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    """Set background color of a table cell (e.g. '003366' or 'F0F4F8')."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cells in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, bold=False, italic=False, font_size=11, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    # Process inline formatting like **bold** and *italic* and `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                label, url = m.groups()
                run = p.add_run(f"{label} ({url})")
                run.font.color.rgb = RGBColor(0, 102, 204)
                run.underline = True
            else:
                p.add_run(token)
        else:
            run = p.add_run(token)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)
            if color:
                run.font.color.rgb = color
    return p

def convert_markdown_to_docx(md_path, docx_path):
    print(f"Mengonversi: {md_path} -> {docx_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Page Setup (Margins: 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 51, 51)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        
        # Parse rows
        rows_data = []
        for line in table_lines:
            line_str = line.strip()
            if not line_str.startswith('|'):
                continue
            # Check separator line like |---|---|
            if re.match(r'^[\|\s\-:]+$', line_str):
                continue
            cells = [c.strip() for c in line_str.split('|')[1:-1]]
            if cells:
                rows_data.append(cells)

        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'

            for r_idx, row_content in enumerate(rows_data):
                row = table.rows[r_idx]
                is_header = (r_idx == 0)
                for c_idx, cell_text in enumerate(row_content):
                    if c_idx >= num_cols:
                        continue
                    cell = row.cells[c_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                    
                    if is_header:
                        set_cell_background(cell, '1F4E79') # Deep blue
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, 'F2F4F7') # Light zebra
                        else:
                            set_cell_background(cell, 'FFFFFF')

                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    
                    # Add formatted text
                    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', cell_text)
                    for token in tokens:
                        if not token:
                            continue
                        if token.startswith('**') and token.endswith('**'):
                            run = p.add_run(token[2:-2])
                            run.bold = True
                        elif token.startswith('*') and token.endswith('*'):
                            run = p.add_run(token[1:-1])
                            run.italic = True
                        elif token.startswith('`') and token.endswith('`'):
                            run = p.add_run(token[1:-1])
                            run.font.name = 'Courier New'
                            run.font.size = Pt(9.5)
                        else:
                            run = p.add_run(token)
                            if is_header:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                run.font.size = Pt(10)

            # Space after table
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(8)

        table_lines = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        
        # Create shaded single cell table for code block
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, 'F4F4F4')
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        run = cp.add_run(code_text.rstrip())
        run.font.name = 'Courier New'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(40, 40, 40)
        
        # Space after code block
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        code_lines = []

    for line in lines:
        raw_line = line
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(46, 117, 182)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(68, 114, 196)
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.bold = True
        elif stripped.startswith('- ') or stripped.startswith('* '):
            add_styled_paragraph(doc, stripped[2:], style='List Bullet', space_after=3)
        elif re.match(r'^\d+\.\s', stripped):
            prefix_len = len(re.match(r'^\d+\.\s', stripped).group(0))
            add_styled_paragraph(doc, stripped[prefix_len:], style='List Number', space_after=3)
        elif stripped.startswith('---') or stripped.startswith('___'):
            # Horizontal rule
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(190, 190, 190)
        else:
            add_styled_paragraph(doc, stripped, style='Normal', space_after=6)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)
    print(f"✅ Selesai: {docx_path}")

def main():
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    md_files = [f for f in os.listdir(docs_dir) if f.endswith('.md')]
    for md_file in md_files:
        md_full = os.path.join(docs_dir, md_file)
        docx_file = md_file.replace('.md', '.docx')
        docx_full = os.path.join(docs_dir, docx_file)
        convert_markdown_to_docx(md_full, docx_full)

if __name__ == "__main__":
    main()
